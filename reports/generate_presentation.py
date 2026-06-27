from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "presentation_churn_telcowave.pptx"


def add_bg(slide, color=(246, 249, 252)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def add_title(slide, title, subtitle=None):
    t = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.8))
    tf = t.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 43, 60)
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.62), Inches(1.05), Inches(11), Inches(0.5))
        sp = s.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = RGBColor(67, 90, 111)


def add_bullets(slide, items, left=0.8, top=1.7, width=5.8, height=4.7):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.font.size = Pt(21)
        p.font.color.rgb = RGBColor(29, 53, 87)
        p.space_after = Pt(8)


def add_kpi_card(slide, title, value, left, top, width=3.8, height=1.5, color=(230, 240, 255)):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*color)
    shape.line.color.rgb = RGBColor(180, 200, 220)
    tf = shape.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(47, 72, 88)
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(17, 43, 60)


def add_img(slide, name, left, top, width):
    path = FIG_DIR / name
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, (239, 246, 255))
    add_title(
        s,
        "Prédiction du Churn Client - TelcoWave",
        "Projet Data Science orienté scoring et priorisation marketing",
    )
    add_bullets(
        s,
        [
            "Objectif: classer les clients par risque de résiliation",
            "Contrainte: budget d'action limité",
            "KPI métier principal: precision@10%",
        ],
        left=0.8,
        top=2.0,
        width=6.4,
        height=3.2,
    )
    add_kpi_card(s, "Taux de churn observé", "~26,5%", left=8.6, top=2.4)

    # 2. Problem & data
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_title(s, "Problème métier, données et pipeline")
    add_bullets(
        s,
        [
            "Dataset Telco Customer Churn (1 ligne = 1 client)",
            "Cible: Churn (Yes/No)",
            "Variables clés: Contract, tenure, PaymentMethod, OnlineSecurity, TechSupport",
            "Preprocessing: imputation TotalCharges + StandardScaler + OneHotEncoder",
            "Split stratifié train / test / validation",
        ],
    )

    # 3. Model comparison visual
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_title(s, "Comparaison des modèles")
    add_img(s, "model_comparison.png", left=0.7, top=1.45, width=12.0)

    # 4. Calibration visual + KPI cards
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_title(s, "Calibration et fiabilité du score")
    add_img(s, "calibration_comparison.png", left=0.7, top=1.5, width=7.2)
    add_kpi_card(s, "Brier score LR calibré", "0,1358", left=8.3, top=2.0, color=(223, 248, 230))
    add_kpi_card(s, "Brier score LR standard", "0,1681", left=8.3, top=3.9, color=(255, 236, 236))

    # 5. Feature importance
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_title(s, "Drivers du churn (importance des variables)")
    add_img(s, "permutation_importance.png", left=0.8, top=1.4, width=11.8)

    # 6. Threshold optimization
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_title(s, "Seuil de ciblage et compromis business")
    add_img(s, "threshold_optimization.png", left=0.7, top=1.45, width=7.5)
    add_bullets(
        s,
        [
            "precision@10% ≈ 0,754",
            "Top 10% clients scorés: forte densité de churners",
            "Permet de concentrer les actions de rétention sur les cas prioritaires",
        ],
        left=8.4,
        top=2.0,
        width=4.4,
        height=3.0,
    )

    # 7. Recommendations
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, (240, 249, 244))
    add_title(s, "Recommandations opérationnelles")
    add_bullets(
        s,
        [
            "Mettre en production la Logistic Regression calibrée",
            "Cibler en priorité les segments month-to-month et faible tenure",
            "Piloter mensuellement: precision@10%, churn évité, ROI campagne",
            "Lancer un test A/B Customer Success pour mesurer l'uplift",
        ],
        left=0.9,
        top=1.8,
        width=11.8,
        height=3.8,
    )

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
