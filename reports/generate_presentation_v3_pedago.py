from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "presentation_churn_pedagogique_v3.pptx"


NAVY = RGBColor(22, 34, 57)
BLUE = RGBColor(41, 98, 255)
LIGHT = RGBColor(247, 250, 255)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(33, 47, 73)
MUTED = RGBColor(96, 109, 131)
GREEN_BG = RGBColor(232, 245, 233)
GREEN = RGBColor(46, 125, 50)


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def header(slide, title, subtitle=None):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.9))
    p = t.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.72), Inches(1.05), Inches(12), Inches(0.5))
        sp = s.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = MUTED


def explain_box(slide, title, body, x, y, w, h):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(218, 226, 240)
    tf = shp.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = BLUE
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK


def bullets(slide, items, x, y, w, h, size=20):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, txt in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)


def img(slide, filename, x, y, w):
    p = FIG_DIR / filename
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w))


def decision_box(slide, text, x=0.8, y=5.9, w=11.8, h=1.0):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = GREEN_BG
    shp.line.color.rgb = RGBColor(199, 230, 201)
    p = shp.text_frame.paragraphs[0]
    p.text = "Conclusion: " + text
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = GREEN


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, LIGHT)
    header(
        s,
        "Prédire le churn client: de l'analyse à la décision",
        "Version pédagogique pour comprendre le raisonnement complet",
    )
    bullets(
        s,
        [
            "Question: quels clients risquent de partir ?",
            "But: agir en priorité sur les clients les plus à risque",
            "Contraintes: budget limité, donc ciblage intelligent nécessaire",
        ],
        0.8,
        1.9,
        7.5,
        2.3,
    )
    explain_box(
        s,
        "Ce que vous allez voir",
        "1) Pourquoi ce problème est important\n2) Comment le modèle a été construit\n3) Quels résultats on obtient\n4) Quelle décision recommander",
        8.6,
        2.0,
        4.0,
        2.7,
    )
    decision_box(
        s,
        "Le modèle sert à prioriser les actions de rétention, pas à remplacer l'équipe métier.",
        y=5.7,
    )

    # Slide 2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, "1. Comprendre le problème métier")
    explain_box(
        s,
        "Pourquoi c'est un enjeu ?",
        "Le churn fait perdre des revenus. Si on ne contacte pas les bons clients, on dépense du budget sans impact.",
        0.8,
        1.5,
        6.0,
        1.8,
    )
    explain_box(
        s,
        "KPI principal",
        "precision@10%: parmi les 10% clients les plus risqués selon le modèle, combien churnent vraiment ?",
        7.0,
        1.5,
        5.5,
        1.8,
    )
    bullets(
        s,
        [
            "Un taux global de churn autour de 26,5% est observé.",
            "Le modèle doit aider à faire mieux qu'un ciblage aléatoire.",
            "On cherche un score exploitable par des non-techniques.",
        ],
        0.9,
        3.7,
        11.8,
        2.0,
        size=19,
    )
    decision_box(s, "Succès = mieux cibler, pas seulement avoir une bonne métrique globale.")

    # Slide 3
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, LIGHT)
    header(s, "2. Comment on a construit le modèle")
    explain_box(
        s,
        "Données utilisées",
        "Ancienneté, type de contrat, moyen de paiement, services activés, charges mensuelles/totales.",
        0.8,
        1.5,
        6.1,
        1.5,
    )
    explain_box(
        s,
        "Préparation des données",
        "Imputation des valeurs manquantes, normalisation numérique, encodage des variables catégorielles.",
        7.0,
        1.5,
        5.5,
        1.5,
    )
    bullets(
        s,
        [
            "Plusieurs modèles ont été comparés: Logistic Regression, Random Forest, XGBoost, LightGBM...",
            "Ensuite, calibration des probabilités pour rendre le score plus fiable.",
            "Les performances sont évaluées sur des données jamais vues.",
        ],
        0.9,
        3.3,
        11.9,
        2.3,
        size=18,
    )
    decision_box(s, "On compare plusieurs options avant de choisir le modèle final.")

    # Slide 4
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, "3. Résultats de comparaison des modèles")
    img(s, "model_comparison.png", 0.75, 1.45, 8.2)
    explain_box(
        s,
        "Comment lire ce graphique ?",
        "Chaque barre compare un modèle sur une métrique. Plus c'est haut (sauf Brier), mieux c'est.",
        9.1,
        1.7,
        3.9,
        1.6,
    )
    explain_box(
        s,
        "Ce qu'on retient",
        "Les modèles sont proches en ROC-AUC, mais la qualité métier dépend du ciblage (precision@10%).",
        9.1,
        3.5,
        3.9,
        1.7,
    )
    decision_box(
        s,
        "La performance pure ne suffit pas: on choisit un modèle utile pour la décision opérationnelle.",
    )

    # Slide 5
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, LIGHT)
    header(s, "4. Pourquoi la calibration change la décision")
    img(s, "calibration_comparison.png", 0.75, 1.45, 7.8)
    explain_box(
        s,
        "Ce que montre la calibration",
        "Quand le modèle dit 70% de risque, on veut que ce soit proche de la réalité.",
        8.8,
        1.8,
        4.1,
        1.6,
    )
    explain_box(
        s,
        "Résultat chiffré",
        "Brier score Logistic calibrée = 0,1358 vs 0,1681 sans calibration.",
        8.8,
        3.6,
        4.1,
        1.6,
    )
    decision_box(
        s,
        "Le modèle calibré est retenu: ses probabilités sont plus fiables pour fixer un seuil de campagne.",
    )

    # Slide 6
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, "5. Ce que l'équipe métier doit faire concrètement")
    bullets(
        s,
        [
            "1) Trier les clients par score de churn (du plus risqué au moins risqué).",
            "2) Cibler d'abord le top 10% (precision@10% ≈ 0,754).",
            "3) Prioriser les profils à risque: faible ancienneté, month-to-month, electronic check, sans services de support/sécurité.",
            "4) Suivre mensuellement: churn évité, coût de campagne, ROI.",
        ],
        0.9,
        1.7,
        12.0,
        3.6,
        size=20,
    )
    img(s, "threshold_optimization.png", 0.85, 4.35, 5.8)
    img(s, "permutation_importance.png", 6.9, 4.35, 5.8)
    decision_box(
        s,
        "Décision finale recommandée: déployer Logistic Regression calibrée et piloter la campagne sur le top clients à risque.",
    )

    # Slide 7
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, LIGHT)
    header(s, "6. Mode d'emploi commercial (très concret)")
    explain_box(
        s,
        "Fichier à utiliser",
        "outputs/ciblage_top10_commercial.csv\n(127 clients déjà triés du plus risqué au moins risqué)",
        0.8,
        1.45,
        6.2,
        1.6,
    )
    explain_box(
        s,
        "Colonnes utiles",
        "customerID, proba_churn, priorite, action_recommandee, Contract, PaymentMethod, tenure...",
        7.2,
        1.45,
        5.2,
        1.6,
    )
    bullets(
        s,
        [
            "Etape 1: ouvrir le fichier et filtrer priorite = 'P1 - Appel immediat'.",
            "Etape 2: appeler dans l'ordre de proba_churn (du plus élevé au plus faible).",
            "Etape 3: utiliser action_recommandee comme trame d'argumentaire.",
            "Etape 4: logger le résultat de l'appel (accepté, refus, rappel) pour mesurer le ROI.",
            "Etape 5: quand les P1 sont traités, passer aux P2 (même logique).",
        ],
        0.9,
        3.35,
        12.0,
        2.7,
        size=18,
    )
    decision_box(
        s,
        "En pratique: le commercial n'a pas besoin du modèle, seulement de la liste priorisée et de l'action recommandée.",
    )

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
