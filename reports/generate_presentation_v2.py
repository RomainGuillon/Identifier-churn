from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "presentation_churn_telcowave_v2.pptx"


NAVY = RGBColor(14, 23, 38)
BLUE = RGBColor(35, 87, 137)
CYAN = RGBColor(51, 168, 203)
LIGHT = RGBColor(245, 248, 252)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(94, 112, 136)
GREEN = RGBColor(40, 167, 69)
RED = RGBColor(220, 53, 69)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def title(slide, text, subtitle=None, dark=False):
    t = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.8), Inches(0.9))
    p = t.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE if dark else NAVY
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.82), Inches(1.15), Inches(11.6), Inches(0.5))
        sp = s.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = RGBColor(214, 227, 243) if dark else MUTED


def bullet_block(slide, items, x, y, w, h, dark=False, size=21):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = it
        p.font.size = Pt(size)
        p.font.color.rgb = WHITE if dark else RGBColor(26, 46, 69)
        p.space_after = Pt(8)


def card(slide, x, y, w, h, bg, label, value, value_color=NAVY):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = bg
    shp.line.color.rgb = RGBColor(220, 228, 238)
    tf = shp.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = MUTED
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = value_color


def image(slide, filename, x, y, w):
    p = FIG_DIR / filename
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w))


def banner(slide, text):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.23)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = CYAN
    shp.line.fill.background()
    t = slide.shapes.add_textbox(Inches(9.4), Inches(0.02), Inches(3.8), Inches(0.2))
    p = t.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, NAVY)
    banner(s, "Projet Data Science - TelcoWave")
    title(
        s,
        "Prédiction du Churn Client",
        "Transformer le scoring en actions de rétention ciblées",
        dark=True,
    )
    bullet_block(
        s,
        [
            "KPI central: precision@10% pour piloter un budget limité",
            "Score exploitable par les équipes Customer Success",
            "Approche: benchmark modèles + calibration des probabilités",
        ],
        0.9,
        2.1,
        7.8,
        3.2,
        dark=True,
        size=22,
    )
    card(s, 9.0, 2.3, 3.4, 1.45, RGBColor(23, 39, 62), "Taux de churn", "26,5%", WHITE)
    card(s, 9.0, 4.0, 3.4, 1.45, RGBColor(23, 39, 62), "Top segment ciblé", "10% clients", WHITE)

    # 2 Story
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, LIGHT)
    banner(s, "Contexte et enjeu")
    title(s, "Le problème métier à résoudre")
    card(s, 0.8, 1.6, 4.0, 2.2, WHITE, "Enjeu", "Réduire la résiliation")
    card(s, 4.95, 1.6, 4.0, 2.2, WHITE, "Contrainte", "Budget marketing limité")
    card(s, 9.1, 1.6, 3.4, 2.2, WHITE, "Décision", "Qui contacter en priorité ?")
    bullet_block(
        s,
        [
            "Un score de risque classe les clients du plus fragile au plus stable.",
            "Les actions de rétention sont focalisées sur la population la plus rentable à traiter.",
            "Le succès se mesure sur la qualité du ciblage, pas seulement sur l'accuracy globale.",
        ],
        0.95,
        4.2,
        11.9,
        2.6,
        size=18,
    )

    # 3 Comparison full
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    banner(s, "Benchmark modèles")
    title(s, "Comparaison de performance des modèles")
    image(s, "model_comparison.png", 0.85, 1.45, 11.6)

    # 4 Calibration
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    banner(s, "Qualité du score")
    title(s, "Calibration: des probabilités plus fiables")
    image(s, "calibration_comparison.png", 0.8, 1.55, 7.6)
    card(s, 8.7, 2.0, 3.8, 1.45, RGBColor(233, 248, 236), "Brier LR calibré", "0,1358", GREEN)
    card(s, 8.7, 3.8, 3.8, 1.45, RGBColor(255, 238, 240), "Brier LR standard", "0,1681", RED)
    bullet_block(
        s,
        [
            "La calibration améliore la confiance dans les probabilités prédites.",
            "Utile pour définir un seuil d'action robuste en production.",
        ],
        8.75,
        5.5,
        3.7,
        1.5,
        size=13,
    )

    # 5 Importance + threshold
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, LIGHT)
    banner(s, "Lecture métier")
    title(s, "Drivers du churn et stratégie de seuil")
    image(s, "permutation_importance.png", 0.7, 1.5, 6.0)
    image(s, "threshold_optimization.png", 6.75, 1.5, 5.9)

    # 6 Action plan
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, NAVY)
    banner(s, "Plan d'action")
    title(s, "Recommandations opérationnelles", dark=True)
    bullet_block(
        s,
        [
            "1. Déployer Logistic Regression calibrée comme moteur de scoring initial.",
            "2. Cibler prioritairement le top 10% clients à risque (precision@10% ≈ 0,754).",
            "3. Lancer un pilote A/B avec Customer Success.",
            "4. Suivre mensuellement: churn évité, ROI, dérive des performances.",
        ],
        0.95,
        1.95,
        12.0,
        4.2,
        dark=True,
        size=23,
    )

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
